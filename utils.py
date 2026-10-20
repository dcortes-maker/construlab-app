"""
Capa de datos — Supabase (PostgreSQL + Storage)
"""
import streamlit as st
from supabase import create_client, Client
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
import io
import json
import secrets

_PROYECTOS_DB = {
    "riviera_park_2":    "riviera-park-2",
    "villas_del_bosque": "villas-del-bosque",
}
BUCKET = "adjuntos"
_sb_cliente = None

def configurar_cliente(client):
    """La API inyecta su cliente. Streamlit lo deja en None y usa secrets."""
    global _sb_cliente
    _sb_cliente = client

def _proyecto_db() -> str:
    pid = st.session_state.get("proyecto", "riviera_park_2")
    return _PROYECTOS_DB.get(pid, "riviera-park-2")
MES      = ["Ene","Feb","Mar","Abr","May","Jun","Jul","Ago","Sep","Oct","Nov","Dic"]
_excluir_matriz = ['separac', 'gasto', 'bonif', '⚠']


def _sb() -> Client:
    if _sb_cliente is not None:
        return _sb_cliente
    cfg = st.secrets["supabase"]
    return create_client(cfg["url"], cfg["service_key"])


def _soltar_cache():
    """Fuera de Streamlit (la API, el respaldo) no hay caché que limpiar."""
    for fn in (cargar_datos, cargar_reservas, clientes_con_contrato, listar_retirados,
               listar_buzon, unidades_con_comprobante):
        try:
            fn.clear()
        except Exception:
            pass


def _a_fecha(v):
    if v is None: return None
    if isinstance(v, date) and not isinstance(v, datetime): return v
    if isinstance(v, datetime): return v.date()
    if isinstance(v, str):
        try: return date.fromisoformat(v[:10])
        except: return None
    return None


def _cat_marca(ref, desc):
    r = str(ref).lower()  if ref  else ''
    d = str(desc).lower() if desc else ''
    if r.startswith('separac'):          return 'separacion'
    if r.startswith('gasto manejo'):     return 'gasto_manejo'
    if r.startswith('gasto legal'):      return 'gasto_legal'
    if r.startswith('bonif'):            return 'bonificacion'
    if r.startswith('abono'):            return 'abono'
    if 'reserva' in d or 'separac' in d: return 'separacion'
    if 'gasto' in d and 'manejo' in d:  return 'gasto_manejo'
    if 'gasto' in d and 'legal'  in d:  return 'gasto_legal'
    if 'bonif' in d:                     return 'bonificacion'
    return 'abono'


@st.cache_data(ttl=15)
def cargar_datos(proyecto: str = ""):
    if not proyecto:
        proyecto = _proyecto_db()
    sb   = _sb()
    rows = sb.table('cuotas').select('*').eq('proyecto', proyecto).order('id').execute().data
    hoy  = date.today()

    vistos  = {}
    sin_plan = set()
    filas   = []

    for row in rows:
        u    = row['unidad']
        nom  = row['nombre']
        desc = row['descripcion'] or ''
        m    = row['monto']
        fv   = _a_fecha(row['fecha_venc'])
        fp   = _a_fecha(row['fecha_pago'])
        ref  = row.get('referencia')

        if u and nom and u not in vistos:
            vistos[u] = nom
            sin_plan.add(u)
        if u and '⚠' in desc:
            sin_plan.add(u)          # marcado explícitamente como sin plan
        elif u and row.get('num_cuota'):
            sin_plan.discard(u)      # tiene cuotas reales → sí tiene plan
        if u and m:
            filas.append({
                'fila':   row['id'],
                'unidad': u,
                'nombre': nom or vistos.get(u, ''),
                'desc':   desc,
                'fv':     fv,
                'monto':  float(m),
                'fp':     fp,
                'ref':    ref,
                'forma':  row.get('forma_pago') or '',
                'num_cuota': row.get('num_cuota'),
            })

    clientes = sorted(vistos.items())

    cuota_map = {}
    for f in filas:
        if not f['fv'] or 'bonif' in f['desc'].lower(): continue
        key = (f['unidad'], f['fv'].year, f['fv'].month)
        if key not in cuota_map:
            cuota_map[key] = {'monto': 0.0, 'pagado': True, 'fecha': f['fv']}
        cuota_map[key]['monto'] += f['monto']
        if not f['fp']:
            cuota_map[key]['pagado'] = False

    all_fechas = [f['fv'] for f in filas
                  if f['fv'] and not any(x in f['desc'].lower() for x in _excluir_matriz)]
    if all_fechas:
        f_min, f_max = min(all_fechas), max(all_fechas)
    else:
        f_min = f_max = date(hoy.year, hoy.month, 1)
    cols = []
    y, m = f_min.year, f_min.month
    while (y, m) <= (f_max.year, f_max.month):
        cols.append((y, m)); m += 1
        if m > 12: m = 1; y += 1

    cols_set      = set(cols)
    pagadas_cnt   = sum(1 for (u,y,m), v in cuota_map.items() if (y,m) in cols_set and v['pagado'])
    atrasadas_cnt = sum(1 for (u,y,m), v in cuota_map.items()
                        if (y,m) in cols_set and not v['pagado']
                        and (v['fecha'].year, v['fecha'].month) < (hoy.year, hoy.month))
    total_cuotas  = sum(1 for (u,y,m) in cuota_map if (y,m) in cols_set)

    dash = defaultdict(lambda: {
        'n_cuotas': 0, 'total_plan': 0.0, 'atrasado': 0.0,
        'separacion': 0.0, 'gasto_manejo': 0.0, 'gasto_legal': 0.0,
        'bonificacion': 0.0, 'abono': 0.0
    })
    for f in filas:
        u = f['unidad']
        if u in sin_plan: continue
        if 'bonif' in f['desc'].lower():
            if f['fp']: dash[u]['bonificacion'] += f['monto']
            continue
        dash[u]['n_cuotas']   += 1
        dash[u]['total_plan'] += f['monto']
        if f['fp']:
            cat = _cat_marca(f['ref'], f['desc'])
            dash[u][cat] += f['monto']
        elif f['fv'] and (f['fv'].year, f['fv'].month) < (hoy.year, hoy.month):
            dash[u]['atrasado'] += f['monto']
    for u in sin_plan:
        for f in filas:
            if f['unidad'] != u: continue
            d = f['desc'].lower()
            if f['fp'] and ('reserva' in d or 'separac' in d):
                dash[u]['separacion'] += f['monto']

    total_abonado  = sum(d['separacion']+d['gasto_manejo']+d['gasto_legal']+d['abono']
                         for d in dash.values())
    total_atrasado = sum(d['atrasado'] for d in dash.values())

    return {
        'clientes': clientes,
        'sin_plan': sin_plan,
        'filas':    filas,
        'cuota_map': cuota_map,
        'cols':     cols,
        'dash':     dict(dash),
        'hoy':      hoy,
        'kpis': {
            'clientes':     len(clientes),
            'total_cuotas': total_cuotas,
            'pagadas':      pagadas_cnt,
            'atrasadas':    atrasadas_cnt,
            'abonado':      total_abonado,
            'atrasado':     total_atrasado,
        }
    }


def marcar_pago(record_id: int, fecha_pago: date, forma_pago: str = 'Transferencia'):
    _sb().table('cuotas').update({
        'fecha_pago': fecha_pago.isoformat(),
        'forma_pago': forma_pago,
        'referencia': f"Pagado - {fecha_pago.strftime('%d/%m/%Y')}",
    }).eq('id', record_id).execute()
    _soltar_cache()


def desmarcar_pago(record_id: int):
    _sb().table('cuotas').update({
        'fecha_pago': None,
        'forma_pago': None,
        'referencia': None,
    }).eq('id', record_id).execute()
    anular_recibo(record_id)
    _soltar_cache()


def pago_parcial(record_id: int, monto_pagado: float, fecha_pago: date,
                 forma_pago: str = 'Transferencia') -> float:
    """Abono menor a la cuota: marca lo pagado y deja el saldo restante como cuota pendiente.

    La fila original queda con el monto abonado y marcada como pagada; se crea una
    fila nueva con el saldo (misma unidad, cuota y fecha de vencimiento).
    Retorna el saldo que quedó pendiente.
    """
    sb  = _sb()
    row = sb.table('cuotas').select('*').eq('id', record_id).execute().data
    if not row:
        return 0.0
    r      = row[0]
    total  = float(r['monto'])
    pagado = round(min(monto_pagado, total), 2)
    saldo  = round(total - pagado, 2)

    sb.table('cuotas').update({
        'monto':      pagado,
        'fecha_pago': fecha_pago.isoformat(),
        'forma_pago': forma_pago,
        'referencia': f"Abono parcial - {fecha_pago.strftime('%d/%m/%Y')}",
    }).eq('id', record_id).execute()

    if saldo > 0:
        sb.table('cuotas').insert({
            'proyecto':    r['proyecto'],
            'unidad':      r['unidad'],
            'nombre':      r['nombre'],
            'num_cuota':   r['num_cuota'],
            'descripcion': f"{r['descripcion']} (saldo)",
            'fecha_venc':  r['fecha_venc'],
            'monto':       saldo,
        }).execute()

    _soltar_cache()
    return saldo


def ajustar_monto_siguiente(record_id: int, excedente: float):
    sb  = _sb()
    row = sb.table('cuotas').select('monto').eq('id', record_id).execute().data
    if row:
        nuevo = max(0, float(row[0]['monto']) - excedente)
        sb.table('cuotas').update({'monto': round(nuevo, 2)}).eq('id', record_id).execute()
    _soltar_cache()


def agregar_filas_plan_bulk(unidad: str, nombre: str, filas: list, proyecto: str = "") -> int:
    sb     = _sb()
    proyecto = proyecto or _proyecto_db()

    existing = sb.table('cuotas').select('num_cuota,descripcion').eq('proyecto', proyecto).eq('unidad', unidad).execute().data
    existentes = {(r['num_cuota'], (r['descripcion'] or '').strip()) for r in existing if r['num_cuota']}

    to_insert = []
    for f in filas:
        clave = (int(f['num_cuota']), str(f['desc']).strip())
        if clave in existentes: continue
        rec = {
            'proyecto': proyecto, 'unidad': unidad, 'nombre': nombre,
            'num_cuota': f['num_cuota'], 'descripcion': f['desc'], 'monto': f['monto'],
        }
        if f.get('fecha_venc'):
            fv = f['fecha_venc']
            rec['fecha_venc'] = fv.isoformat() if hasattr(fv, 'isoformat') else str(fv)
        to_insert.append(rec)
        existentes.add(clave)

    if to_insert:
        for i in range(0, len(to_insert), 100):
            sb.table('cuotas').insert(to_insert[i:i+100]).execute()
    _soltar_cache()
    return len(to_insert)


def eliminar_filas_plan(record_ids: list):
    sb = _sb()
    for rid in record_ids:
        sb.table('cuotas').delete().eq('id', rid).execute()
    _soltar_cache()


def actualizar_fila_plan(record_id: int, desc: str, fecha_venc, monto: float):
    _sb().table('cuotas').update({
        'descripcion': desc,
        'fecha_venc':  fecha_venc.isoformat() if fecha_venc and hasattr(fecha_venc, 'isoformat') else None,
        'monto':       monto,
    }).eq('id', record_id).execute()
    _soltar_cache()


def eliminar_plan(unidad: str):
    _sb().table('cuotas').delete().eq('proyecto', _proyecto_db()).eq('unidad', unidad).execute()
    _soltar_cache()


def agregar_fila_plan(unidad: str, nombre: str, num_cuota, desc: str, fecha_venc, monto: float):
    """Inserta una cuota y retorna el id creado (o None si falla)."""
    rec = {
        'proyecto': _proyecto_db(), 'unidad': unidad, 'nombre': nombre,
        'num_cuota': num_cuota, 'descripcion': desc, 'monto': monto,
    }
    if fecha_venc:
        rec['fecha_venc'] = fecha_venc.isoformat() if hasattr(fecha_venc, 'isoformat') else str(fecha_venc)
    res = _sb().table('cuotas').insert(rec).execute()
    _soltar_cache()
    return res.data[0]['id'] if res.data else None


def reordenar_plan(ids_en_orden: list):
    """Renumera num_cuota = 1..N siguiendo el orden de los ids recibidos."""
    sb = _sb()
    for pos, cid in enumerate(ids_en_orden, 1):
        sb.table('cuotas').update({'num_cuota': pos}).eq('id', int(cid)).execute()
    _soltar_cache()


# ─── Reservas ──────────────────────────────────────────────────────

@st.cache_data(ttl=15)
def cargar_reservas(proyecto: str = "") -> list:
    if not proyecto:
        proyecto = _proyecto_db()
    return _sb().table('reservas').select('*').eq('proyecto', proyecto).order('id').execute().data


def guardar_reserva(nombre: str, unidad: str, monto: float, fecha, notas: str = '') -> int:
    res = _sb().table('reservas').insert({
        'proyecto': _proyecto_db(),
        'nombre':   nombre.strip().upper(),
        'unidad':   unidad.strip().upper(),
        'monto':    float(monto),
        'fecha':    fecha.isoformat() if hasattr(fecha, 'isoformat') else str(fecha),
        'notas':    notas.strip(),
    }).execute()
    return res.data[0]['id']


def migrar_reserva_a_plan(unidad: str, proyecto: str = "") -> dict:
    """Al cargar el plan de un cliente en reserva: marca la cuota "Reserva" como
    pagada con la fecha y el monto de la reserva, y borra el registro de reservas
    para que el dinero no quede contado dos veces.

    Devuelve {} si la unidad no tenía reserva.
    """
    sb  = _sb()
    proyecto = proyecto or _proyecto_db()
    res = sb.table('reservas').select('*') \
            .eq('proyecto', proyecto).eq('unidad', unidad).execute().data
    if not res:
        return {}
    r          = res[0]
    monto_res  = float(r['monto'])
    fecha_res  = _a_fecha(r['fecha']) or date.today()

    # Cuota de reserva del plan (la primera cuyo concepto sea "Reserva")
    cuotas = sb.table('cuotas').select('id,num_cuota,descripcion,monto,fecha_pago') \
               .eq('proyecto', proyecto).eq('unidad', unidad) \
               .order('num_cuota').execute().data
    # "Reserva" o "Separación": ambos nombran el primer abono de reserva
    cuota = next((c for c in cuotas
                  if not c['fecha_pago']
                  and any(k in (c['descripcion'] or '').lower()
                          for k in ('reserva', 'separac'))), None)

    marcada = False
    if cuota:
        if monto_res >= float(cuota['monto']):
            marcar_pago(cuota['id'], fecha_res, 'Transferencia')
        else:
            pago_parcial(cuota['id'], monto_res, fecha_res, 'Transferencia')
        marcada = True

    sb.table('reservas').delete().eq('id', r['id']).execute()
    _soltar_cache()
    return {'monto': monto_res, 'fecha': fecha_res, 'marcada': marcada}


def eliminar_reserva(rid: int):
    _sb().table('reservas').delete().eq('id', rid).execute()


# ─── Retiros (quien se sale del proyecto) ─────────────────────────
# Retirar no borra nada. A las filas del cliente se les cambia el proyecto por
# "<proyecto>#retirado#<id>", así desaparecen de todas las consultas de la app
# —que siempre filtran por proyecto— y el apartamento queda libre para otro
# cliente sin que los dos historiales se mezclen. Deshacerlo es devolverle su
# proyecto original.
#
# La ficha de cada retiro vive en su propio archivo del bucket. Un archivo por
# retiro y no una lista compartida: así dos personas retirando a la vez no se
# pisan lo que escribió la otra.

_MARCA_RETIRO = '#retirado#'
_DIR_RETIROS  = 'sistema/retiros'


def _proyecto_retirado(proyecto: str, rid: str) -> str:
    return f"{proyecto}{_MARCA_RETIRO}{rid}"


def _leer_doc(ruta: str):
    """Lee un JSON del bucket sin pasar por la caché de Supabase.

    storage.download() se sirve cacheado y devolvería la versión anterior
    durante una hora, así que se pide por URL firmada con un parámetro
    distinto en cada llamada.
    """
    import time as _time
    import urllib.request as _url
    try:
        firmada = _sb().storage.from_(BUCKET).create_signed_url(ruta, 60)["signedURL"]
    except Exception:
        return None
    sep = "&" if "?" in firmada else "?"
    ultimo = None
    for intento in range(3):
        pedido = _url.Request(f"{firmada}{sep}cb={_time.time_ns()}",
                              headers={"Cache-Control": "no-cache"})
        try:
            with _url.urlopen(pedido, timeout=20) as resp:
                return json.loads(resp.read().decode())
        except Exception as e:   # la red se cae a ratos; reintentar es más barato que fallar
            ultimo = e
            _time.sleep(1 + intento)
    raise ultimo


def _escribir_doc(ruta: str, data: dict):
    _sb().storage.from_(BUCKET).upload(
        ruta, json.dumps(data, ensure_ascii=False, default=str).encode(),
        file_options={"content-type": "application/json",
                      "upsert": "true", "cache-control": "0"},
    )


def _ficha_retiro(proyecto, unidad, nombre, tipo, motivo, abonado,
                  aplica_devolucion, usuario) -> dict:
    rid = f"{datetime.now().strftime('%Y%m%d%H%M%S')}{secrets.token_hex(3)}"
    ficha = {
        'id':           rid,
        'proyecto':     proyecto,
        'unidad':       unidad,
        'nombre':       nombre or '',
        'tipo':         tipo,
        'fecha_retiro': date.today().isoformat(),
        'motivo':       (motivo or '').strip(),
        'abonado':      round(float(abonado), 2),
        'devolucion':   'pendiente' if aplica_devolucion else 'no_aplica',
        'devol_monto':  None,
        'devol_fecha':  None,
        'creado_por':   usuario or '',
        'created_at':   datetime.now().isoformat(timespec='seconds'),
    }
    _escribir_doc(f"{_DIR_RETIROS}/{rid}.json", ficha)
    return ficha


def retirar_cliente(unidad: str, motivo: str = '', aplica_devolucion: bool = False,
                    usuario: str = '', proyecto: str = '') -> dict:
    sb = _sb()
    proyecto = proyecto or _proyecto_db()
    filas = (sb.table('cuotas').select('id,nombre,monto,fecha_pago')
             .eq('proyecto', proyecto).eq('unidad', unidad).execute().data) or []
    if not filas:
        return {'ok': False, 'motivo': 'Ese apartamento no tiene un plan activo.'}

    abonado = sum(float(f['monto'] or 0) for f in filas if f.get('fecha_pago'))
    nombre  = next((f['nombre'] for f in filas if f.get('nombre')), '')

    ficha = _ficha_retiro(proyecto, unidad, nombre, 'cpp', motivo,
                          abonado, aplica_devolucion, usuario)
    sb.table('cuotas').update({'proyecto': _proyecto_retirado(proyecto, ficha['id'])}) \
      .in_('id', [f['id'] for f in filas]).execute()
    _soltar_cache()
    return {'ok': True, 'retiro_id': ficha['id'], 'nombre': nombre,
            'abonado': round(abonado, 2), 'cuotas': len(filas)}


def retirar_reserva(rid_reserva: int, motivo: str = '', aplica_devolucion: bool = False,
                    usuario: str = '') -> dict:
    sb = _sb()
    filas = sb.table('reservas').select('*').eq('id', rid_reserva).execute().data or []
    if not filas:
        return {'ok': False, 'motivo': 'Esa reserva ya no existe.'}
    r = filas[0]
    if _MARCA_RETIRO in (r.get('proyecto') or ''):
        return {'ok': False, 'motivo': 'Esa reserva ya está retirada.'}

    ficha = _ficha_retiro(r['proyecto'], r['unidad'], r.get('nombre') or '', 'reserva',
                          motivo, float(r.get('monto') or 0), aplica_devolucion, usuario)
    sb.table('reservas').update({'proyecto': _proyecto_retirado(r['proyecto'], ficha['id'])}) \
      .eq('id', rid_reserva).execute()
    _soltar_cache()
    return {'ok': True, 'retiro_id': ficha['id'], 'nombre': r.get('nombre') or '',
            'abonado': round(float(r.get('monto') or 0), 2)}


@st.cache_data(ttl=15)
def listar_retirados(proyecto: str = "") -> list:
    proyecto = proyecto or _proyecto_db()
    try:
        archivos = _sb().storage.from_(BUCKET).list(
            _DIR_RETIROS, {'limit': 1000}) or []
    except Exception:
        return []
    fichas = []
    for a in archivos:
        nombre = a.get('name') or ''
        if not nombre.endswith('.json'):
            continue
        ficha = _leer_doc(f"{_DIR_RETIROS}/{nombre}")
        if ficha and ficha.get('proyecto') == proyecto:
            fichas.append(ficha)
    return sorted(fichas, key=lambda f: f.get('created_at') or '', reverse=True)


def marcar_devolucion_hecha(retiro_id: str, monto: float = None, fecha=None) -> bool:
    ruta = f"{_DIR_RETIROS}/{retiro_id}.json"
    ficha = _leer_doc(ruta)
    if not ficha:
        return False
    ficha['devolucion']  = 'devuelta'
    ficha['devol_fecha'] = (fecha or date.today()).isoformat()
    if monto is not None:
        ficha['devol_monto'] = round(float(monto), 2)
    _escribir_doc(ruta, ficha)
    _soltar_cache()
    return True


def reingresar_retirado(retiro_id: str) -> dict:
    """Deshace un retiro: el cliente vuelve a aparecer como si nada.

    Se niega si el apartamento ya se le vendió a otro: devolver al anterior
    dejaría a dos clientes distintos bajo el mismo número y los totales de esa
    unidad saldrían sumados.
    """
    sb = _sb()
    ficha = _leer_doc(f"{_DIR_RETIROS}/{retiro_id}.json")
    if not ficha:
        return {'ok': False, 'motivo': 'No encuentro ese retiro.'}

    ocupada = (sb.table('cuotas').select('nombre')
               .eq('proyecto', ficha['proyecto']).eq('unidad', ficha['unidad'])
               .limit(1).execute().data) or []
    if not ocupada:
        ocupada = (sb.table('reservas').select('nombre')
                   .eq('proyecto', ficha['proyecto']).eq('unidad', ficha['unidad'])
                   .limit(1).execute().data) or []
    if ocupada:
        quien = ocupada[0].get('nombre') or 'otro cliente'
        return {'ok': False,
                'motivo': f"El apartamento {ficha['unidad']} ya es de {quien}. "
                          f"Para reingresar a {ficha['nombre']} hay que retirar antes a esa persona."}

    marcado = _proyecto_retirado(ficha['proyecto'], retiro_id)
    sb.table('cuotas').update({'proyecto': ficha['proyecto']}).eq('proyecto', marcado).execute()
    sb.table('reservas').update({'proyecto': ficha['proyecto']}).eq('proyecto', marcado).execute()
    try:
        sb.storage.from_(BUCKET).remove([f"{_DIR_RETIROS}/{retiro_id}.json"])
    except Exception:
        pass
    _soltar_cache()
    return {'ok': True, 'nombre': ficha['nombre'], 'unidad': ficha['unidad']}


# ─── Buzón (lo que se encontró en Drive y falta aprobar) ──────────
# Lo llena scripts/revisar_drive.py. Una propuesta no crea nada: alguien la
# aprueba desde la pantalla Buzón y ahí sí entra el plan.

_DIR_BUZON = 'sistema/buzon'


@st.cache_data(ttl=15)
def listar_buzon(proyecto: str = "", solo_pendientes: bool = True) -> list:
    proyecto = proyecto or _proyecto_db()
    try:
        archivos = _sb().storage.from_(BUCKET).list(_DIR_BUZON, {'limit': 1000}) or []
    except Exception:
        return []
    out = []
    for a in archivos:
        nombre = a.get('name') or ''
        # Los que empiezan por "_" son del sistema, como el índice del detector.
        if not nombre.endswith('.json') or nombre.startswith('_'):
            continue
        try:
            p = _leer_doc(f"{_DIR_BUZON}/{nombre}")
        except Exception:
            continue   # una propuesta ilegible no puede tumbar la pantalla entera
        if not p or p.get('proyecto') != proyecto:
            continue
        if solo_pendientes and p.get('estado') != 'pendiente':
            continue
        out.append(p)
    return sorted(out, key=lambda p: (not p.get('legible'), p.get('unidad') or ''))


def _cerrar_propuesta(pid: str, estado: str, quien: str = ''):
    ruta = f"{_DIR_BUZON}/{pid}.json"
    p = _leer_doc(ruta)
    if not p:
        return None
    p['estado'] = estado
    p['resuelto_por'] = quien or ''
    p['resuelto_en'] = datetime.now().isoformat(timespec='seconds')
    _escribir_doc(ruta, p)
    _soltar_cache()
    return p


def aprobar_propuesta(pid: str, unidad: str = '', nombre: str = '', quien: str = '') -> dict:
    """Mete al cliente en la app: su documento y, si se pudo leer, su plan.

    Sin cuotas el cliente igual entra, marcado «Plan de pagos pendiente» —así
    figura ya en la app con su contrato guardado, y el plan se carga después.
    """
    p = _leer_doc(f"{_DIR_BUZON}/{pid}.json")
    if not p:
        return {'ok': False, 'motivo': 'No encuentro esa propuesta.'}

    unidad = (unidad or p['unidad']).strip().upper()
    nombre = (nombre or p['nombre']).strip().upper()

    if p.get('cuotas'):
        filas = [{'num_cuota': c['num_cuota'], 'desc': c['desc'],
                  'fecha_venc': c.get('fecha_venc') or None, 'monto': c['monto']}
                 for c in p['cuotas']]
        creadas = agregar_filas_plan_bulk(unidad, nombre, filas, p['proyecto'])
    else:
        creadas = agregar_filas_plan_bulk(
            unidad, nombre,
            [{'num_cuota': 1, 'desc': '⚠ Plan de pagos pendiente', 'monto': 0.0}],
            p['proyecto'])

    # El documento pasa del buzón a la carpeta del cliente, con un nombre que
    # clasificar_documento reconozca para que el respaldo nocturno lo copie.
    guardado = ''
    if p.get('adjunto'):
        try:
            ext = p['adjunto'].rsplit('.', 1)[-1]
            prefijo = 'Contrato' if p.get('tipo') == 'contrato' else 'Plan_de_Pagos_anexo'
            guardado = f"{prefijo}_{unidad.replace('/', '-')}.{ext}"
            datos = _sb().storage.from_(BUCKET).download(p['adjunto'])
            subir_adjunto(unidad, guardado, datos)
            _sb().storage.from_(BUCKET).remove([p['adjunto']])
        except Exception:
            guardado = ''   # el cliente ya entró; el documento se puede subir a mano

    _cerrar_propuesta(pid, 'aprobada', quien)
    return {'ok': True, 'unidad': unidad, 'nombre': nombre,
            'creadas': creadas, 'con_plan': bool(p.get('cuotas')), 'documento': guardado}


def descartar_propuesta(pid: str, quien: str = '') -> bool:
    p = _leer_doc(f"{_DIR_BUZON}/{pid}.json")
    if p and p.get('adjunto'):
        try:
            _sb().storage.from_(BUCKET).remove([p['adjunto']])
        except Exception:
            pass
    return bool(_cerrar_propuesta(pid, 'descartada', quien))


def saldos_retiros(proyecto: str = "") -> dict:
    """Qué pasó con el dinero de los que se fueron.

    Sin devolución el abono se queda en el proyecto, así que sigue sumando al
    total en banco aunque su plan de pagos ya no exista. Con devolución sale:
    pendiente mientras no se haga la transferencia, devuelto una vez hecha.
    """
    retenido = por_devolver = devuelto = 0.0
    n_ret = n_pend = 0
    for r in listar_retirados(proyecto):
        monto = float(r.get('abonado') or 0)
        estado = r.get('devolucion')
        if estado == 'no_aplica':
            retenido += monto
            n_ret += 1 if monto else 0
        elif estado == 'pendiente':
            por_devolver += monto
            n_pend += 1
        else:
            devuelto += float(r.get('devol_monto') or monto)
    return {'retenido': round(retenido, 2), 'cuantos_retenido': n_ret,
            'por_devolver': round(por_devolver, 2), 'cuantas_pendientes': n_pend,
            'devuelto': round(devuelto, 2)}


# ─── Adjuntos (Supabase Storage) ──────────────────────────────────

def _folder(unidad: str) -> str:
    return unidad.replace('/', '-').replace(' ', '_')


def _mime(filename: str) -> str:
    ext = filename.lower().rsplit('.', 1)[-1]
    return {
        'pdf':  'application/pdf',
        'png':  'image/png',
        'jpg':  'image/jpeg',
        'jpeg': 'image/jpeg',
        'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'xls':  'application/vnd.ms-excel',
        'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'doc':  'application/msword',
    }.get(ext, 'application/octet-stream')


def listar_adjuntos(unidad: str) -> list:
    try:
        folder = _folder(unidad)
        files  = _sb().storage.from_(BUCKET).list(folder)
        return [f for f in (files or []) if f.get('name') and not f['name'].startswith('.')]
    except Exception:
        return []


@st.cache_data(ttl=60)
def clientes_con_contrato(proyecto: str = "") -> set:
    """Unidades con un adjunto cuyo nombre contiene 'contrato' (en Plan de Pagos → Documentos)."""
    out = set()
    try:
        datos = cargar_datos(proyecto)
        for u, _ in datos['clientes']:
            for f in listar_adjuntos(u):
                nombre = f['name'].lower()
                if 'contrato' in nombre or 'cpp' in nombre:
                    out.add(u)
                    break
    except Exception:
        pass
    return out


@st.cache_data(ttl=15)
def unidades_con_comprobante(unidades: tuple) -> set:
    """Cuáles de esas unidades ya tienen guardado un comprobante de pago."""
    out = set()
    for u in unidades:
        for f in listar_adjuntos(u):
            if clasificar_documento(f['name']) == 'comprobante':
                out.add(u)
                break
    return out


def subir_adjunto(unidad: str, filename: str, file_bytes: bytes) -> bool:
    sb     = _sb()
    folder = _folder(unidad)
    path   = f"{folder}/{filename}"
    subir_adjunto.ultimo_error = ""
    try:
        sb.storage.from_(BUCKET).remove([path])
    except Exception:
        pass
    try:
        sb.storage.from_(BUCKET).upload(
            path, file_bytes,
            file_options={"content-type": _mime(filename), "upsert": "true"}
        )
        return True
    except Exception as e:
        subir_adjunto.ultimo_error = str(e)
        try:
            st.error(f"Error al subir archivo: {e}")
        except Exception:
            pass
        return False


def descargar_adjunto(unidad: str, filename: str) -> bytes:
    return _sb().storage.from_(BUCKET).download(f"{_folder(unidad)}/{filename}")


def eliminar_adjunto(unidad: str, filename: str):
    _sb().storage.from_(BUCKET).remove([f"{_folder(unidad)}/{filename}"])


def clasificar_documento(filename: str) -> str:
    """contrato, anexo, comprobante, o vacío si no entra en el respaldo de Drive."""
    n = (filename or "").lower()
    if "contrato" in n or "cpp" in n:
        return "contrato"
    if n.startswith("plan_de_pagos") or "anexo" in n or "plan de pago" in n:
        return "anexo"
    if "comprobante" in n:
        return "comprobante"
    return ""


def listar_unidades(proyecto: str) -> list:
    """Unidades con plan o reserva. Sirve a la API, sin la caché de Streamlit."""
    sb = _sb()
    vistos = {}
    rows = sb.table('cuotas').select('unidad,nombre').eq('proyecto', proyecto).limit(5000).execute().data or []
    for r in rows:
        if r.get('unidad') and r['unidad'] not in vistos:
            vistos[r['unidad']] = r.get('nombre') or ''
    reservas = sb.table('reservas').select('unidad,nombre').eq('proyecto', proyecto).limit(1000).execute().data or []
    for r in reservas:
        if r.get('unidad'):
            vistos.setdefault(r['unidad'], r.get('nombre') or '')
    return [{'unidad': u, 'nombre': n} for u, n in sorted(vistos.items())]


def ruta_pdf_recibo(num: int, unidad: str) -> str:
    return f"recibos/{_folder(unidad)}/Recibo_{num}.pdf"


def guardar_pdf_recibo(num: int, unidad: str, pdf_bytes: bytes) -> bool:
    """Deja el PDF tal cual se emitió. El respaldo copia este archivo."""
    sb   = _sb()
    path = ruta_pdf_recibo(num, unidad)
    try:
        sb.storage.from_(BUCKET).remove([path])
    except Exception:
        pass
    try:
        sb.storage.from_(BUCKET).upload(
            path, pdf_bytes,
            file_options={"content-type": "application/pdf", "upsert": "true"}
        )
        return True
    except Exception:
        return False


def leer_pdf_recibo(num: int, unidad: str):
    try:
        return _sb().storage.from_(BUCKET).download(ruta_pdf_recibo(num, unidad))
    except Exception:
        return None


def obtener_pdf_recibo(num: int, unidad: str, nombre: str, desc: str, monto: float,
                       fecha_pago, forma_pago: str = 'Transferencia') -> bytes:
    """Devuelve el PDF guardado. Si todavía no existe, lo genera y lo guarda."""
    ya = leer_pdf_recibo(num, unidad)
    if ya:
        return ya
    pdf = generar_recibo(nombre, unidad, desc, monto, fecha_pago, num, forma_pago)
    guardar_pdf_recibo(num, unidad, pdf)
    return pdf


# ─── Generación de PDFs ────────────────────────────────────────────

def pdf_a_imagenes(pdf_bytes: bytes, scale: float = 2.0) -> list:
    """Convierte cada página de un PDF a PNG (para vista previa en la nube)."""
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(pdf_bytes)
    paginas = []
    for page in doc:
        img = page.render(scale=scale).to_pil()
        b = io.BytesIO()
        img.save(b, format='PNG')
        paginas.append(b.getvalue())
    return paginas

def siguiente_num_recibo() -> int:
    """Contador global consecutivo de recibos (inicia en 1020), guardado en Supabase."""
    sb = _sb()
    try:
        row = sb.table('contadores').select('valor').eq('nombre', 'recibo').execute().data
        if row:
            nuevo = row[0]['valor'] + 1
            sb.table('contadores').update({'valor': nuevo}).eq('nombre', 'recibo').execute()
        else:
            nuevo = 1020
            sb.table('contadores').insert({'nombre': 'recibo', 'valor': nuevo}).execute()
        return nuevo
    except Exception:
        return 1020


def registrar_recibo(num: int, cuota_id: int, unidad: str, nombre: str,
                     desc: str, monto: float, fecha_pago, forma: str = ''):
    """Guarda el recibo emitido en Supabase (para trazabilidad y anulación)."""
    try:
        _sb().table('recibos').insert({
            'num':      num,
            'cuota_id': cuota_id,
            'unidad':   unidad,
            'nombre':   nombre,
            'desc':     desc,
            'monto':    monto,
            'fecha':    fecha_pago.isoformat() if hasattr(fecha_pago, 'isoformat') else str(fecha_pago),
            'estado':   'emitido',
            'proyecto': _proyecto_db(),
            'forma':    forma,
        }).execute()
    except Exception:
        pass


def listar_recibos() -> list:
    """Recibos del proyecto actual, más recientes primero."""
    try:
        return _sb().table('recibos').select('*') \
                    .eq('proyecto', _proyecto_db()) \
                    .order('num', desc=True).execute().data or []
    except Exception:
        return []


def anular_recibo(cuota_id: int):
    """Marca como ANULADO el recibo emitido de una cuota desmarcada (el número no se reutiliza)."""
    try:
        _sb().table('recibos').update({'estado': 'anulado'}) \
             .eq('cuota_id', cuota_id).eq('estado', 'emitido').execute()
    except Exception:
        pass


def generar_recibo(nombre: str, unidad: str, desc: str, monto: float,
                   fecha_pago, num_recibo: int, forma_pago: str = 'Transferencia') -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas as _canvas
    from reportlab.lib import colors

    buf = io.BytesIO()
    W, H = letter
    c = _canvas.Canvas(buf, pagesize=letter)

    LOGO   = str(Path(__file__).parent / "logo_recibo.png")
    MARCA  = str(Path(__file__).parent / "logo_recibo_marca.png")
    BLACK  = colors.black
    RED    = colors.HexColor('#C00000')

    # ── Logo superior izquierdo ────────────────────────────────────
    logo_w = 230
    logo_h = logo_w * 683 / 1024
    c.drawImage(LOGO, 75, H - 40 - logo_h, width=logo_w, height=logo_h,
                preserveAspectRatio=True, mask='auto')

    # ── Datos de la empresa (centrado bajo el logo) ────────────────
    cx_emp = 75 + logo_w / 2
    c.setFillColor(BLACK); c.setFont("Helvetica", 10)
    c.drawCentredString(cx_emp, H - 183, "Riviera Park Development S.A.")
    c.drawCentredString(cx_emp, H - 199, "Las Mañanitas, Ciudad de Panamá")
    c.drawCentredString(cx_emp, H - 214, "+507 6975-4414 / 375-0404")
    c.drawCentredString(cx_emp, H - 230, "ventas@gterrabona.com")

    # ── Título ─────────────────────────────────────────────────────
    c.setFont("Helvetica-Bold", 20)
    c.drawString(349, H - 150, "RECIBO DE PAGO")

    # ── Secciones CLIENTE / RECIBO ─────────────────────────────────
    c.setFont("Helvetica", 12)
    c.drawString(90, H - 283, "CLIENTE")
    c.drawString(306, H - 283, "RECIBO")

    fecha_str = fecha_pago.strftime('%-d/%-m/%Y') if hasattr(fecha_pago, 'strftime') else str(fecha_pago)
    metodo = {'transferencia': 'ACH', 'efectivo': 'Efectivo', 'cheque': 'Cheque'}.get(
        forma_pago.lower(), forma_pago)

    # El nombre se envuelve en varias líneas para no invadir la columna RECIBO (x=306).
    c.setFont("Helvetica", 12)
    ANCHO_NOMBRE = 306 - 90 - 12          # margen de seguridad antes del N° de recibo
    etiqueta     = "Nombre: "
    disponible   = ANCHO_NOMBRE - c.stringWidth(etiqueta, "Helvetica", 12)

    lineas, actual = [], ""
    for palabra in str(nombre).split():
        prueba = f"{actual} {palabra}".strip()
        ancho  = c.stringWidth(prueba, "Helvetica", 12)
        limite = disponible if not lineas else ANCHO_NOMBRE
        if actual and ancho > limite:
            lineas.append(actual)
            actual = palabra
        else:
            actual = prueba
    if actual:
        lineas.append(actual)

    yn = H - 323
    for i, ln in enumerate(lineas):
        c.drawString(90, yn, f"{etiqueta}{ln}" if i == 0 else ln)
        yn -= 16

    y = H - 323
    c.drawString(306, y, "No.: ")
    c.setFillColor(RED); c.setFont("Helvetica-Bold", 12)
    c.drawString(306 + c.stringWidth("No.: ", "Helvetica", 12), y, str(num_recibo))
    c.setFillColor(BLACK); c.setFont("Helvetica", 12)
    y -= 19.5; c.drawString(306, y, f"Fecha: {fecha_str}")
    y -= 19.5; c.drawString(306, y, f"Método de pago: {metodo}")
    y -= 19.5; c.drawString(306, y, "Estado: Pagado")
    y -= 19.5; c.drawString(306, y, "Moneda: USD")

    # ── Tabla Descripción / Cantidad / Total ───────────────────────
    tx0, tx1, tx2, tx3 = 85, 343, 428, 559
    th_top = H - 443          # borde superior del header
    th_bot = H - 473          # borde inferior del header
    tr_bot = H - 501          # borde inferior de la fila de datos

    c.setLineWidth(1); c.setStrokeColor(BLACK)
    for yy in (th_top, th_bot, tr_bot):
        c.line(tx0, yy, tx3, yy)
    for xx in (tx0, tx1, tx2, tx3):
        c.line(xx, th_top, xx, tr_bot)

    c.setFont("Helvetica-Bold", 12)
    c.drawString(tx0 + 5, th_top - 20, "Descripción")
    c.drawString(tx1 + 6, th_top - 20, "Cantidad")
    c.drawString(tx2 + 6, th_top - 20, "Total")

    c.setFont("Helvetica", 12)
    c.drawString(tx0 + 5, th_bot - 19, f"{desc} — Unidad {unidad}"[:48])
    c.drawString(tx1 + 6, th_bot - 19, "1")
    c.drawString(tx2 + 9, th_bot - 19, f"${monto:,.2f}")

    # ── Observaciones ──────────────────────────────────────────────
    c.setFont("Helvetica", 11)
    c.drawString(90, H - 564, "Observaciones:")
    c.drawString(90, H - 595, "_" * 42)

    # ── Marca de agua inferior derecha ─────────────────────────────
    wm_w = 354
    wm_h = wm_w * 683 / 1024
    c.drawImage(MARCA, 221, 20, width=wm_w, height=wm_h,
                preserveAspectRatio=True, mask='auto')

    c.save()
    buf.seek(0)
    return buf.read()


def generar_estado_cuenta(unidad: str, nombre: str, filas: list) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import cm, mm
    from reportlab.pdfgen import canvas as _canvas
    from reportlab.lib import colors
    from datetime import date as _date

    buf = io.BytesIO()
    W, H = letter
    c = _canvas.Canvas(buf, pagesize=letter)

    TEAL     = colors.HexColor('#0D5C6E')
    TEAL_LT  = colors.HexColor('#E6F4F7')
    GREEN    = colors.HexColor('#1B5E20')
    GREEN_LT = colors.HexColor('#E8F5E9')
    RED      = colors.HexColor('#B71C1C')
    RED_LT   = colors.HexColor('#FFEBEE')
    GRAY     = colors.HexColor('#555555')
    LGRAY    = colors.HexColor('#F7F7F7')
    WHITE    = colors.white
    BLACK    = colors.black
    LOGO_PATH = str(Path(__file__).parent / "logo_recibo.png")

    hoy   = _date.today()
    M     = 1.8 * cm
    avail = W - 2*M

    filas_s      = sorted(filas, key=lambda x: (x['fv'] or _date(2099,1,1), x['desc']))
    total_plan   = sum(f['monto'] for f in filas_s)
    total_pagado = sum(f['monto'] for f in filas_s if f['fp'])
    saldo        = total_plan - total_pagado

    COLS = [0.9*cm, 5.8*cm, 3.4*cm, 3.2*cm, 4.1*cm]
    HDRS = ['#', 'DESCRIPCI\xd3N', 'F. VENCIMIENTO', 'MONTO', 'ESTADO']
    COLS[1] += avail - sum(COLS)

    def draw_header():
        hdr_h = 2.8*cm
        logo_w = 4.2*cm
        logo_h = logo_w * 683 / 1024
        c.drawImage(LOGO_PATH, M, H - M - logo_h - 0.1*cm,
                    width=logo_w, height=logo_h, preserveAspectRatio=True, mask='auto')
        c.setFillColor(BLACK); c.setFont("Helvetica-Bold", 15)
        c.drawRightString(W - M, H - M - 0.9*cm, "ESTADO DE CUENTA")
        c.setFillColor(GRAY); c.setFont("Helvetica", 9)
        c.drawRightString(W - M, H - M - 1.5*cm, "Riviera Park Development S.A.")
        c.drawRightString(W - M, H - M - 1.9*cm, f"Fecha: {hoy.strftime('%d/%m/%Y')}")
        c.setStrokeColor(TEAL); c.setLineWidth(1.5)
        c.line(M, H - M - hdr_h, W - M, H - M - hdr_h)
        return H - M - hdr_h

    def draw_table_header(y):
        c.setFillColor(TEAL)
        c.rect(M, y - 0.6*cm, avail, 0.6*cm, fill=1, stroke=0)
        c.setFillColor(WHITE); c.setFont("Helvetica-Bold", 7.5)
        cx = M
        for hdr, cw in zip(HDRS, COLS):
            c.drawString(cx + 2*mm, y - 0.41*cm, hdr); cx += cw
        return y - 0.6*cm

    y = draw_header()
    y -= 0.3*cm
    c.setFillColor(TEAL_LT)
    c.rect(M, y - 1.3*cm, avail, 1.3*cm, fill=1, stroke=0)
    c.setFillColor(TEAL); c.setFont("Helvetica-Bold", 8.5)
    c.drawString(M + 0.3*cm, y - 0.52*cm, "CLIENTE:")
    c.setFillColor(BLACK); c.setFont("Helvetica-Bold", 10)
    c.drawString(M + 2.2*cm, y - 0.52*cm, nombre.upper())
    c.setFillColor(TEAL); c.setFont("Helvetica-Bold", 8.5)
    c.drawString(M + 0.3*cm, y - 1.03*cm, "UNIDAD:")
    c.setFillColor(BLACK); c.setFont("Helvetica", 10)
    c.drawString(M + 2.2*cm, y - 1.03*cm, unidad)
    y -= 1.3*cm

    y -= 0.4*cm
    bw3 = (avail - 0.4*cm) / 3

    def sum_box(x, label, val, bg, col):
        c.setFillColor(bg)
        c.roundRect(x, y - 1.5*cm, bw3, 1.5*cm, 3*mm, fill=1, stroke=0)
        c.setStrokeColor(col); c.setLineWidth(0.5)
        c.roundRect(x, y - 1.5*cm, bw3, 1.5*cm, 3*mm, fill=0, stroke=1)
        c.setFillColor(GRAY); c.setFont("Helvetica", 7.5)
        c.drawCentredString(x + bw3/2, y - 0.55*cm, label)
        c.setFillColor(col); c.setFont("Helvetica-Bold", 14)
        c.drawCentredString(x + bw3/2, y - 1.25*cm, f"$ {val:,.2f}")

    sum_box(M,                   "TOTAL DEL PLAN",  total_plan,   TEAL_LT, TEAL)
    sum_box(M + bw3 + 0.2*cm,   "TOTAL PAGADO",    total_pagado, GREEN_LT, GREEN)
    sum_box(M + 2*(bw3+0.2*cm), "SALDO PENDIENTE", saldo,
            RED_LT if saldo > 0 else GREEN_LT, RED if saldo > 0 else GREEN)
    y -= 1.5*cm

    y -= 0.4*cm
    y  = draw_table_header(y)
    RH = 0.55*cm

    for i, f in enumerate(filas_s):
        if y - RH < M + 1.5*cm:
            _draw_footer(c, M, avail, W, hoy)
            c.showPage()
            y = draw_header()
            y = draw_table_header(y)

        bg = (GREEN_LT if f['fp']
              else RED_LT if f['fv'] and (f['fv'].year, f['fv'].month) < (hoy.year, hoy.month)
              else (WHITE if i % 2 == 0 else LGRAY))
        c.setFillColor(bg)
        c.rect(M, y - RH, avail, RH, fill=1, stroke=0)

        fv_str = f['fv'].strftime('%d/%m/%Y') if f['fv'] else '—'
        if f['fp']:
            estado_txt, estado_col = f"Pag. {f['fp'].strftime('%d/%m/%Y')}", GREEN
        elif f['fv'] and (f['fv'].year, f['fv'].month) < (hoy.year, hoy.month):
            estado_txt, estado_col = 'ATRASADO', RED
        else:
            estado_txt, estado_col = 'Pendiente', GRAY

        c.setFillColor(BLACK); c.setFont("Helvetica", 7.5)
        cx = M
        for val, cw in zip([str(i+1), f['desc'], fv_str, f"$ {f['monto']:,.2f}"], COLS[:-1]):
            c.drawString(cx + 2*mm, y - 0.38*cm, val); cx += cw
        c.setFillColor(estado_col); c.setFont("Helvetica-Bold", 7)
        c.drawString(cx + 2*mm, y - 0.38*cm, estado_txt)

        c.setStrokeColor(colors.HexColor('#E0E0E0')); c.setLineWidth(0.3)
        c.line(M, y - RH, M + avail, y - RH)
        y -= RH

    c.setFillColor(TEAL_LT); c.setStrokeColor(TEAL); c.setLineWidth(1)
    c.rect(M, y - RH - 1*mm, avail, RH + 1*mm, fill=1, stroke=1)
    c.setFillColor(TEAL); c.setFont("Helvetica-Bold", 9)
    c.drawString(M + 0.3*cm, y - 0.38*cm, "TOTAL DEL PLAN")
    c.drawString(M + sum(COLS[:3]) + 2*mm, y - 0.38*cm, f"$ {total_plan:,.2f}")

    _draw_footer(c, M, avail, W, hoy)
    c.save()
    buf.seek(0)
    return buf.read()


def _draw_footer(c, M, avail, W, hoy):
    from reportlab.lib import colors
    TEAL = colors.HexColor('#0D5C6E')
    from reportlab.lib.units import cm, mm
    c.setFillColor(TEAL)
    c.rect(M, M, avail, 0.55*cm, fill=1, stroke=0)
    c.setFillColor(colors.white); c.setFont("Helvetica", 6.5)
    c.drawCentredString(W/2, M + 1.5*mm,
        f"Riviera Park Development, S.A.  —  Proyecto Riviera Park II  —  Generado el {hoy.strftime('%d/%m/%Y')}")


def extraer_datos_cliente(archivo_bytes, nombre_archivo: str) -> dict:
    """Lee nombre y unidad del ANEXO (ambos formatos). Devuelve {} si no los halla."""
    import re
    if not str(nombre_archivo).lower().endswith('.pdf'):
        return {}
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(archivo_bytes)) as pdf:
            texto = '\n'.join((pg.extract_text() or '') for pg in pdf.pages)
    except Exception:
        return {}
    if not texto.strip():
        return {}

    out = {}

    # Unidad: "Unidad 4-H – T1 – F1" (anexo nuevo) o "Unidad: 3-C" (anexo de tabla)
    mu = re.search(r'Unidad:?\s*([A-Za-z]{0,2}\d{0,2}[A-Za-z]?[-–]?[A-Za-z0-9]+)', texto)
    if mu:
        out['unidad'] = mu.group(1).strip().upper().replace('–', '-')

    # Nombre del comprador, en sus distintas variantes
    for patron in (r'Comprador\s*1\s+(.+?)\s+(?:Tel[eé]fono|C[eé]dula|Email)',
                   r'Persona a notificar:\s*(.+)',
                   r'Persona a\s+(.+?)\s*\n\s*notificar:'):
        mn = re.search(patron, texto)
        if mn:
            nom = re.sub(r'\s+', ' ', mn.group(1)).strip(' :').upper()
            if nom:
                out['nombre'] = nom
            break

    return out


# ─── parsear_plan — sin cambios (procesa archivos subidos) ─────────

def parsear_plan(archivo_bytes, nombre_archivo: str, _debug=False) -> list:
    import re
    from datetime import date as date_cls

    MESES = {"ene":1,"feb":2,"mar":3,"abr":4,"may":5,"jun":6,
              "jul":7,"ago":8,"sep":9,"oct":10,"nov":11,"dic":12,
              "jan":1,"apr":4,"aug":8,"dec":12}

    def _ultimo_dia(anio, mes):
        import calendar
        return calendar.monthrange(anio, mes)[1]

    def _parse_fecha(txt):
        if txt is None: return None
        if isinstance(txt, (date_cls, datetime)):
            d = txt.date() if isinstance(txt, datetime) else txt
            return d.replace(day=_ultimo_dia(d.year, d.month))
        txt = str(txt).strip().lower()
        m = re.search(r'([a-z]{3})[- ](\d{2,4})', txt)
        if m:
            mes = MESES.get(m.group(1))
            anio = int(m.group(2))
            if anio < 100: anio += 2000
            if mes: return date_cls(anio, mes, _ultimo_dia(anio, mes))
        m2 = re.search(r'(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})', txt)
        if m2:
            d2, mo, yr = int(m2.group(1)), int(m2.group(2)), int(m2.group(3))
            if yr < 100: yr += 2000
            try: return date_cls(yr, mo, _ultimo_dia(yr, mo))
            except: pass
        return None

    def _parse_monto(txt):
        if txt is None: return None
        if isinstance(txt, (int, float)): return float(txt) if txt > 0 else None
        s = str(txt).strip().replace('$','').replace(' ','')
        s = re.sub(r'(?i)b/\.?', '', s)          # símbolo balboa
        s = re.sub(r'[^\d.,]', '', s)

        if ',' in s and '.' in s:
            # El separador que aparece de último es el decimal
            if s.rfind(',') > s.rfind('.'):
                s = s.replace('.', '').replace(',', '.')   # 1.050,25
            else:
                s = s.replace(',', '')                     # 1,050.25
        elif ',' in s:
            # Una sola coma: es decimal solo si le siguen 1 o 2 dígitos.
            # "1,050" son mil cincuenta, no 1.05.
            ent, _, dec = s.rpartition(',')
            s = f"{ent.replace(',', '')}.{dec}" if len(dec) in (1, 2) else s.replace(',', '')
        elif s.count('.') > 1:
            s = s.replace('.', '')                         # 1.050.000

        s = re.sub(r'[^\d.]', '', s)
        try:
            v = float(s); return v if v > 0 else None
        except: return None

    def _normalizar_desc(desc_raw):
        d = str(desc_raw).lower()
        if 'reserva' in d:                    return 'Reserva'
        if 'firma' in d or 'separac' in d:    return 'Separación'
        if 'gasto' in d and 'manejo' in d:    return 'Gasto Manejo'
        if 'gasto' in d and 'legal' in d:     return 'Gasto Legal'
        if 'bonif' in d:                      return 'Bonificación'
        return desc_raw.strip() or 'Abono inicial'

    def _procesar_filas(rows):
        col_num = col_fecha = col_ab = col_ex = col_desc = None
        resultado = []; en_tabla = False
        for row in rows:
            if not row: continue
            cells = [str(c).strip() if c is not None else '' for c in row]
            low   = [c.lower() for c in cells]
            if not en_tabla:
                def _col(*claves):
                    # Entre los encabezados que coinciden, elige el más corto: evita
                    # las celdas fusionadas que arrastran toda la columna de datos.
                    cands = [i for i, c in enumerate(low) if any(k in c for k in claves)]
                    return min(cands, key=lambda i: len(low[i])) if cands else None

                if any('cuot' in c for c in low) and any('fech' in c for c in low):
                    col_num  = _col('cuot')
                    col_fecha= _col('fech')
                    col_ab   = _col('abono', 'inicial')
                    col_ex   = _col('gasto', 'manejo', 'legal')
                    col_desc = _col('observ', 'concepto', 'desc')
                    en_tabla = True
                continue
            if col_num is None or col_num >= len(cells): continue
            num_raw = cells[col_num]
            if not re.match(r'^\d+$', num_raw): continue
            num = int(num_raw)
            fecha    = _parse_fecha(row[col_fecha] if col_fecha is not None and col_fecha < len(row) else None)
            monto_ab = _parse_monto(row[col_ab]   if col_ab   is not None and col_ab   < len(row) else None)
            monto_ex = _parse_monto(row[col_ex]   if col_ex   is not None and col_ex   < len(row) else None)
            desc_raw = cells[col_desc] if col_desc is not None and col_desc < len(cells) else ''
            if not desc_raw:
                for c in cells:
                    if c and not re.match(r'^[\d$.,/-]+$', c) and not re.match(r'^\d+$', c):
                        desc_raw = c; break
            desc    = _normalizar_desc(desc_raw)
            desc_ex = 'Gasto Legal' if col_ex is not None and 'legal' in (low[col_ex] if col_ex < len(low) else '') else 'Gasto Manejo'
            desc_ab = desc if 'gasto' not in desc.lower() else 'Abono inicial'
            if monto_ab and monto_ab > 0:
                resultado.append({'num_cuota': num, 'desc': desc_ab, 'fecha_venc': fecha, 'monto': monto_ab})
            if monto_ex and monto_ex > 0:
                resultado.append({'num_cuota': num, 'desc': desc_ex, 'fecha_venc': fecha, 'monto': monto_ex})
        return resultado

    ext = nombre_archivo.lower().rsplit('.', 1)[-1]

    if ext in ('xlsx', 'xls'):
        import openpyxl, pandas as pd
        wb_tmp = openpyxl.load_workbook(io.BytesIO(archivo_bytes), data_only=True)
        fls = []
        for sheet in wb_tmp.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row = list(row)
                if all(v is None or str(v).strip() == '' for v in row): continue
                fls.append(row)
        resultado = _procesar_filas(fls)
        if resultado: return resultado
        df_tmp = pd.read_excel(io.BytesIO(archivo_bytes), header=None)
        return _procesar_filas(df_tmp.values.tolist())

    ROMANOS = ['I','II','III','IV','V','VI','VII','VIII','IX','X',
               'XI','XII','XIII','XIV','XV','XVI','XVII','XVIII','XIX','XX']

    def _fecha_mmdd(txt):
        """Fecha del ANEXO en formato MM/DD/YYYY, respetando el día exacto."""
        m = re.match(r'(\d{1,2})/(\d{1,2})/(\d{4})', str(txt).strip())
        if not m:
            return None
        a, b, anio = int(m.group(1)), int(m.group(2)), int(m.group(3))
        mes, dia = (a, b) if a <= 12 else (b, a)   # si el 1° > 12 es DD/MM
        try:
            return date_cls(anio, mes, min(dia, _ultimo_dia(anio, mes)))
        except Exception:
            return None

    def _parsear_anexo(texto):
        """Lee el ANEXO de Riviera Park (formato etiqueta-valor, no tabla).

        Toma solo la sección PLAN DE PAGOS: reserva, Pago N y gastos.
        """
        if 'PLAN DE PAGOS' not in texto.upper():
            return []
        res = []

        m = re.search(r'Fecha de reserva\s+(\d{1,2}/\d{1,2}/\d{4}).*?'
                      r'Monto de reserva\s+B/\.?\s*([\d.,]+)', texto, re.I)
        if m:
            monto = _parse_monto(m.group(2))
            if monto:
                res.append({'num_cuota': 1, 'desc': 'Reserva',
                            'fecha_venc': _fecha_mmdd(m.group(1)), 'monto': monto})

        for mp in re.finditer(r'Fecha Pago\s+(\d+)\s+(\d{1,2}/\d{1,2}/\d{4})\s+'
                              r'Monto Pago\s+\1\s+B/\.?\s*([\d.,]+)', texto, re.I):
            num   = int(mp.group(1))
            monto = _parse_monto(mp.group(3))
            if not monto:
                continue
            desc = f"{ROMANOS[num-1]} Pago" if 1 <= num <= len(ROMANOS) else 'Abono inicial'
            res.append({'num_cuota': num, 'desc': desc,
                        'fecha_venc': _fecha_mmdd(mp.group(2)), 'monto': monto})

        sig = max((r['num_cuota'] for r in res), default=0) + 1
        for patron, etiqueta in ((r'Gastos?\s+legales?\s+B/\.?\s*([\d.,]+)', 'Gasto Legal'),
                                 (r'Gastos?\s+de\s+Manejo\s+B/\.?\s*([\d.,]+)', 'Gasto Manejo')):
            mg = re.search(patron, texto, re.I)
            if mg:
                monto = _parse_monto(mg.group(1))
                if monto:
                    res.append({'num_cuota': sig, 'desc': etiqueta,
                                'fecha_venc': None, 'monto': monto})
                    sig += 1

        res.sort(key=lambda r: r['num_cuota'])
        return res

    import pdfplumber
    fls, texto_pdf = [], ''
    with pdfplumber.open(io.BytesIO(archivo_bytes)) as pdf:
        total_chars = sum(len(p.chars) for p in pdf.pages)
        if total_chars > 0:
            for page in pdf.pages:
                for tabla in page.extract_tables():
                    fls.extend(tabla)
                texto_pdf += (page.extract_text() or '') + '\n'
            resultado = _procesar_filas(fls)
            if resultado: return resultado
            resultado = _parsear_anexo(texto_pdf)
            if resultado: return resultado

    return []

